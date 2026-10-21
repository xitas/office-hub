"""PDF and Excel files from a report Document (see documents.py).

PDF (fpdf2): A4 landscape with the fonts embedded (subset): Noto Sans for Latin text and Noto Nastaliq Urdu for
Urdu, shaped by HarfBuzz so letters join and stack properly, with the Unicode bidi algorithm for mixed text. An Urdu
report reads right to left: header text on the right and table columns in mirrored order.

Excel (openpyxl): one sheet per table under the report's title and header lines; numbers, percentages and dates as
real values; Urdu sheets right to left. Text cells are always stored as text, never as formulas, and those that start
like a formula (= + - @, tab, CR) are marked as quoted text, the same rule as the CSV exports.
"""
import io
import math
import re
from pathlib import Path

from django.utils.translation import gettext as _
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.contacts.csv_io import FORMULA_CHARS

from .documents import Document, fmt_date

FONTS = Path(__file__).resolve().parent / "fonts"


def _text(value, kind) -> str:
    if value is None or value == "":
        return "—" if kind in ("pct", "date") else ""
    if kind == "pct":
        return f"{math.floor(value * 100 + 0.5)}%"  # half up, as on screen (Python's round() would give 12.5 → 12)
    if kind == "date":
        return fmt_date(value)
    return str(value)


# ================================================================ PDF


class _PDF(FPDF):
    def __init__(self, rtl: bool):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.rtl = rtl
        self.add_font("Sans", "", str(FONTS / "NotoSans-Regular.ttf"))
        self.add_font("Sans", "B", str(FONTS / "NotoSans-Bold.ttf"))
        self.add_font("Urdu", "", str(FONTS / "NotoNastaliqUrdu-Regular.ttf"))
        self.add_font("Urdu", "B", str(FONTS / "NotoNastaliqUrdu-Bold.ttf"))
        self.main = "Urdu" if rtl else "Sans"
        self.set_fallback_fonts(["Sans" if rtl else "Urdu"], exact_match=False)
        self.shaping(True)
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(True, margin=14)
        self.alias_nb_pages()

    def shaping(self, on: bool):
        """HarfBuzz shaping and bidi: Urdu reports read right to left; English ones follow each paragraph's own
        direction (names may be in Urdu). Off for text with no right-to-left letters (much faster)."""
        if on:
            self.set_text_shaping(use_shaping_engine=True, direction="rtl" if self.rtl else None)
        else:
            self.set_text_shaping(False)

    def footer(self):
        self.set_y(-10)
        self.set_font(self.main, size=7)
        self.set_text_color(120)
        self.cell(0, 6, _("Page %(n)s of %(total)s") % {"n": self.page_no(), "total": "{nb}"}, align="C")
        self.set_text_color(0)

    @property
    def line_factor(self) -> float:
        return 2.3 if self.rtl else 1.5  # Nastaliq is tall: more room per line


PAD = 1.2  # mm inside each cell
RTL_CHARS = re.compile("[\u0590-\u08ff\ufb1d-\ufdff\ufe70-\ufeff]")  # Hebrew, Arabic and Urdu letters
HEADER_FILL, STRIPE_FILL, BORDER = (235, 238, 242), (248, 249, 251), (190, 196, 204)


def _draw_table(pdf: _PDF, columns, rows: list[list[str]], rtl: bool):
    """A plain table: header (repeated on each new page), striped rows, cells wrapped only when they don't fit.
    Drawn directly rather than with fpdf2's table(), which lays every cell out twice and took most of an export's
    time (about 0.4 s for 60 rows); most cells are short numbers that fit on one line."""
    total = sum(c.width for c in columns)
    widths = [pdf.epw * c.width / total for c in columns]
    aligns = [("R" if rtl else "L") if c.kind == "text" else "C" for c in columns]
    line_h = 8 * 0.35 * pdf.line_factor
    pdf.set_draw_color(*BORDER)
    pdf.set_line_width(0.2)

    measured: dict = {}  # the same text in the same column and style (statuses, priorities…) is measured once

    def lines_of(text: str, width: float) -> list[str]:
        key = (text, width, pdf.font_style)
        if key not in measured:
            if not text or pdf.get_string_width(text) <= width - 2 * PAD:
                measured[key] = [text]
            else:
                measured[key] = pdf.multi_cell(width - 2 * PAD, line_h, text, dry_run=True, output="LINES") or [text]
        return measured[key]

    def draw_row(cells, fill, bold, header=False):
        pdf.set_font(pdf.main, "B" if bold else "", 8)
        # Text with no right-to-left letters (numbers, Latin names) needs neither shaping nor the bidi algorithm,
        # which cost most of the time per cell: those cells are measured and drawn with shaping off.
        plain = [not RTL_CHARS.search(t) for t in cells]
        wrapped = [None] * len(cells)
        for shaped in (False, True):
            pdf.shaping(shaped)
            for i, (text, width) in enumerate(zip(cells, widths)):
                if plain[i] != shaped:
                    wrapped[i] = lines_of(text, width)
        height = max(len(lines) for lines in wrapped) * line_h + 1
        if not header and pdf.get_y() + height > pdf.page_break_trigger:
            pdf.add_page()
            draw_row([c.label for c in columns], HEADER_FILL, True, header=True)
            pdf.set_font(pdf.main, "B" if bold else "", 8)
        y, x = pdf.get_y(), pdf.l_margin
        lefts = []
        for width in widths:
            if fill:
                pdf.set_fill_color(*fill)
            pdf.rect(x, y, width, height, style="DF" if fill else "D")
            lefts.append(x)
            x += width
        for shaped in (False, True):
            pdf.shaping(shaped)
            for i, (lines, width, align, left) in enumerate(zip(wrapped, widths, aligns, lefts)):
                if plain[i] == shaped:
                    continue
                for n, line in enumerate(lines):
                    pdf.set_xy(left + PAD, y + 0.5 + n * line_h)
                    pdf.cell(width - 2 * PAD, line_h, line, align=align)
        pdf.set_xy(pdf.l_margin, y + height)

    if pdf.get_y() + 3 * line_h > pdf.page_break_trigger:
        pdf.add_page()
    draw_row([c.label for c in columns], HEADER_FILL, True, header=True)
    for i, cells in enumerate(rows):
        draw_row(cells, STRIPE_FILL if i % 2 else None, False)
    if not rows:
        pdf.set_font(pdf.main, "", 8)
        pdf.cell(pdf.epw, line_h + 1, _("Nothing to show."), border=1, align="C", new_x="LMARGIN", new_y="NEXT")


def render_pdf(doc: Document) -> bytes:
    pdf = _PDF(doc.rtl)
    pdf.set_title(doc.title)
    pdf.set_creator("Office CRM")
    pdf.add_page()
    side = "R" if doc.rtl else "L"
    pdf.set_font(pdf.main, "B", 15)
    pdf.multi_cell(0, 15 * 0.35 * pdf.line_factor, doc.title, align=side, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(pdf.main, "", 9)
    for line in doc.lines:
        pdf.multi_cell(0, 9 * 0.35 * pdf.line_factor, line, align=side, new_x="LMARGIN", new_y="NEXT")
    for table in doc.tables:
        pdf.ln(3)
        pdf.set_font(pdf.main, "B", 11)
        pdf.multi_cell(0, 11 * 0.35 * pdf.line_factor, table.title, align=side, new_x="LMARGIN", new_y="NEXT")
        columns = list(reversed(table.columns)) if doc.rtl else table.columns
        rows = [[_text(v, c.kind) for v, c in zip(values, table.columns)] for values in table.rows]
        _draw_table(pdf, columns, [list(reversed(r)) if doc.rtl else r for r in rows], doc.rtl)
        if table.note:
            pdf.set_font(pdf.main, "", 8)
            pdf.multi_cell(0, 8 * 0.35 * pdf.line_factor, table.note, align=side, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ================================================================ Excel


def _sheet_name(title: str, used: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:31] or "Sheet"
    base, n = name, 2
    while name in used:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(name)
    return name


def _excel_date_format() -> str:
    """The office's date format, in Excel's terms."""
    from apps.core.models import OrganizationSettings

    return {"MM/DD/YYYY": "mm/dd/yyyy", "YYYY-MM-DD": "yyyy-mm-dd", "DD MMM YYYY": "dd mmm yyyy"}.get(
        OrganizationSettings.cached().date_format, "dd/mm/yyyy")


def _put(ws, row, col, value, kind="text", **style):
    cell = ws.cell(row=row, column=col)
    if value is None or value == "":
        cell.value = None
    elif kind in ("int", "pct", "date"):
        cell.value = value
        cell.number_format = {"int": "0", "pct": "0%", "date": _excel_date_format()}[kind]
    else:
        text = str(value)
        cell.value = text
        cell.data_type = "s"  # always text, never a formula
        if text.startswith(FORMULA_CHARS):
            cell.quotePrefix = True
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def render_xlsx(doc: Document) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Office CRM"
    wb.properties.title = doc.title
    used: set = set()
    header_fill = PatternFill("solid", fgColor="EBEEF2")
    side = "right" if doc.rtl else "left"
    for table in doc.tables:
        ws = wb.create_sheet(_sheet_name(table.title, used))
        ws.sheet_view.rightToLeft = doc.rtl
        _put(ws, 1, 1, doc.title, font=Font(bold=True, size=14), alignment=Alignment(horizontal=side))
        for i, line in enumerate(doc.lines, start=2):
            _put(ws, i, 1, line, alignment=Alignment(horizontal=side))
        top = len(doc.lines) + 3
        _put(ws, top, 1, table.title, font=Font(bold=True, size=12), alignment=Alignment(horizontal=side))
        for j, column in enumerate(table.columns, start=1):
            _put(ws, top + 1, j, column.label, font=Font(bold=True), fill=header_fill, alignment=Alignment(wrap_text=True, vertical="top"))
            ws.column_dimensions[get_column_letter(j)].width = max(10, round(column.width * 12))
        for i, values in enumerate(table.rows, start=top + 2):
            for j, (value, column) in enumerate(zip(values, table.columns), start=1):
                _put(ws, i, j, value, column.kind)
        if table.note:
            _put(ws, top + 2 + len(table.rows), 1, table.note, font=Font(italic=True))
        ws.freeze_panes = ws.cell(row=top + 2, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
