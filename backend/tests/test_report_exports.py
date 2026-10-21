"""Reports, slice 5: every report's PDF and Excel export, read back, shows the same numbers as the screen.

One dataset touches all seven reports (tasks, client interactions, leads, field visits, errands, attendance and leave)
for a manager's department; for each report and language the "by person" rows of the API response (what the screen
shows) are found again in the PDF's text (a row is a line; in Urdu its cells read in mirrored order) and in the Excel
sheet (real numbers, percentages as fractions, dates as dates). Also: file names (an ASCII one, and the report's title
and dates in the reader's language), the audit log, and that staff can't export."""
import io
import math
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from apps.accounts.models import User
from apps.contacts.models import Contact
from apps.core.models import AuditLog
from apps.core.office_time import office_today, office_tz
from apps.leave.models import LeaveType
from apps.reports.documents import fmt_date
from apps.visits.models import VisitCheck, VisitReport
from tests.test_attendance_reports import check_in, leave
from tests.test_client_reports import _lead, log
from tests.test_field_reports import errand, visit
from tests.test_reports import task

pytestmark = pytest.mark.django_db

BASE = "/api/v1/reports"


@pytest.fixture
def people(make_user, admin, manager, staff):
    out = {"admin": admin, "manager": manager, "staff": staff, "colleague": make_user("colleague@test.com", full_name="Sara Colleague")}
    User.objects.update(date_joined=datetime(2020, 1, 1, tzinfo=office_tz()))
    return out


def at(day, hh, mm=0):
    return datetime.combine(day, time(hh, mm), tzinfo=office_tz())


def working_days_back(n):
    today, out = office_today(), []
    day = today - timedelta(days=1)
    while len(out) < n:
        if day.weekday() < 5:
            out.append(day)
        day -= timedelta(days=1)
    return out


@pytest.fixture
def world(people):
    today = office_today()
    s, c = people["staff"], people["colleague"]
    # tasks: on time, late, overdue (high), and one for the colleague
    task("On time", [s], today - timedelta(days=5), done_at=at(today - timedelta(days=6), 12))
    task("Late", [s], today - timedelta(days=3), done_at=at(today - timedelta(days=1), 12))
    task("Overdue", [s], today - timedelta(days=10), priority="high")
    task("Hers", [c], today - timedelta(days=2))
    # client interactions and leads
    client = Contact.objects.create(first_name="Client", assigned_to=s)
    log("call", s, at(today - timedelta(days=4), 10), contact=client)
    log("meeting", c, at(today - timedelta(days=3), 11), contact=client)
    _lead("Lead one", s, at(today - timedelta(days=8), 10), ("contacted", at(today - timedelta(days=6), 10)), ("won", at(today - timedelta(days=2), 10)))
    _lead("Lead two", c, at(today - timedelta(days=7), 10), ("lost", at(today - timedelta(days=3), 10)))
    # field visits and errands
    done = visit(s, -3, "done")
    VisitCheck.objects.create(visit=done, user=s, checked_in_at=done.start, in_location="declined")
    VisitReport.objects.create(visit=done, outcome="good", written_by=s)
    visit(c, -2, "cancelled")
    errand(s, -2, "done", finished_delta=timedelta(hours=-1))
    errand(c, -1, "failed", finished_delta=timedelta(hours=1))
    # attendance and leave
    days = working_days_back(4)
    check_in(s, days[0], 9, 40)
    check_in(s, days[1], 9, mode="remote")
    check_in(c, days[0], 11, mode="field")
    annual = LeaveType.objects.create(name="Annual", name_ur="سالانہ", days_per_year=Decimal("20"), position=1)
    leave(c, annual, days[2], days[2], 1)
    leave(c, annual, days[1], days[1], Decimal("0.5"), half_day="morning")
    return {"from": (today - timedelta(days=20)).isoformat(), "to": today.isoformat()}


def pct_text(v):
    return "—" if v is None else f"{math.floor(v * 100 + 0.5)}%"


def as_text(value, kind):
    if kind == "pct":
        return pct_text(value)
    if kind == "date":
        return "—" if value is None else fmt_date(date.fromisoformat(value))
    if kind == "num":
        return f"{value:g}"
    return str(value)


def as_cell(value, kind):
    return date.fromisoformat(value) if kind == "date" and value else value


# Each report: its URL, extra parameters, which table of the export holds its "by person" rows, and the cells of a
# row from the API response as (value, kind).
SPECS = {
    "task-completion": (True, 2, lambda d: [
        [(p["name"], "text"), (p["department"], "text"), *[(p[k], "int") for k in ("due", "done", "on_time", "late", "overdue", "open")],
         (p["rate"], "pct"), (p["on_time_rate"], "pct"), (p["completed"], "int")] for p in d["people"]]),
    "overdue-tasks": (False, 0, lambda d: [
        [(p["name"], "text"), (p["department"], "text"), *[(p[k], "int") for k in ("overdue", "days_1_7", "days_8_30", "days_31_plus", "urgent")],
         (p["oldest"], "date")] for p in d["people"]]),
    "client-interactions": (True, 1, lambda d: [
        [(p["name"], "text"), (p["department"], "text"), *[(p["counts"][k], "int") for k in d["kinds"]], (p["total"], "int"), (p["clients"], "int")]
        for p in d["people"]]),
    "lead-pipeline": (True, 1, lambda d: [
        [(p["name"], "text"), *[(p[k], "int") for k in ("created", "contacted", "in_discussion", "won", "lost", "open")],
         (p["win_rate"], "pct"), (p["conversion"], "pct")] for p in d["people"] if p["name"]]),
    "field-visits": (True, 1, lambda d: [
        [(p["name"], "text"), *[(p[k], "int") for k in ("scheduled", "done", "cancelled", "not_closed", "upcoming")], (p["completion"], "pct"),
         (p["checked_in"], "int"), (p["reported"], "int")] for p in d["people"]]),
    "errands": (True, 1, lambda d: [
        [(p["name"], "text"), *[(p[k], "int") for k in ("asked", "done", "failed", "cancelled", "open", "overdue")], (p["success"], "pct"),
         (p["on_time_rate"], "pct")] for p in d["people"] if p["name"]]),
    "attendance": (True, 2, lambda d: [
        [(p["name"], "text"), *[(p[k], "int") for k in ("working_days", "office", "field", "remote", "late", "leave", "absent")], (p["rate"], "pct")]
        for p in d["people"]]),
}


def pdf_lines(content) -> set[str]:
    return {" ".join(line.split()) for page in PdfReader(io.BytesIO(content)).pages for line in page.extract_text().split("\n")}


def sheet_rows(ws) -> dict:
    """Data rows of an exported sheet by their first cell (title, header lines, table title and column headers above)."""
    header = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1) if i > 1 and row[0] is not None and row[1] is not None)
    return {row[0]: row for row in ws.iter_rows(min_row=header + 1, values_only=True) if row[0] is not None}


@pytest.mark.parametrize("report", list(SPECS))
def test_exports_show_the_screen_numbers(client_for, people, world, report):
    with_period, table, rows_of = SPECS[report]
    params = world if with_period else {}
    manager = client_for(people["manager"])
    for lang in ("en", "ur"):
        data = manager.get(f"{BASE}/{report}/", params, HTTP_ACCEPT_LANGUAGE=lang).json()
        expected = rows_of(data)
        assert {r[0][0] for r in expected} >= {"Staff", "Sara Colleague"}, report
        assert any(v not in (0, None) for r in expected for v, k in r[1:] if k != "text"), f"{report}: no data to compare"

        pdf = manager.get(f"{BASE}/{report}/export/", {**params, "file": "pdf"}, HTTP_ACCEPT_LANGUAGE=lang)
        assert pdf.status_code == 200 and pdf["Content-Disposition"].startswith(f'attachment; filename="{report.replace("attendance", "attendance-leave")}')
        local = unquote(pdf["Content-Disposition"].split("filename*=UTF-8''")[1])
        assert local.endswith(".pdf") and (params.get("to", office_today().isoformat()) in local)
        assert bool(re.search(r"[\u0600-\u06ff]", local)) == (lang == "ur"), local  # the title in the reader's language
        lines = pdf_lines(pdf.content)
        for row in expected:
            cells = [as_text(v, k) for v, k in row if v not in ("", None) or k != "text"]
            line = " ".join(reversed(cells) if lang == "ur" else cells)
            assert line in lines, f"{report} ({lang}) PDF: {line!r} not found"

        xlsx = manager.get(f"{BASE}/{report}/export/", {**params, "file": "xlsx"}, HTTP_ACCEPT_LANGUAGE=lang)
        wb = load_workbook(io.BytesIO(xlsx.content))
        ws = wb.worksheets[table]
        assert ws.sheet_view.rightToLeft == (lang == "ur")
        got = sheet_rows(ws)
        for row in expected:
            values = tuple(as_cell(v, k) for v, k in row)
            cells = got[row[0][0]][: len(values)]
            cells = tuple(c.date() if isinstance(c, datetime) else c for c in cells)
            assert cells == tuple(None if v == "" else v for v in values), f"{report} ({lang}) Excel: {cells} != {values}"
    assert AuditLog.objects.filter(action="export").count() == 4
    assert client_for(people["staff"]).get(f"{BASE}/{report}/export/", {**params, "file": "pdf"}).status_code == 403


def test_attendance_leave_table_matches(client_for, people, world):
    manager = client_for(people["manager"])
    data = manager.get(f"{BASE}/attendance/", world).json()
    types = data["leave_types"]
    for lang in ("en", "ur"):
        lines = pdf_lines(manager.get(f"{BASE}/attendance/export/", {**world, "file": "pdf"}, HTTP_ACCEPT_LANGUAGE=lang).content)
        ws = load_workbook(io.BytesIO(manager.get(f"{BASE}/attendance/export/", {**world, "file": "xlsx"}, HTTP_ACCEPT_LANGUAGE=lang).content)).worksheets[3]
        got = sheet_rows(ws)
        for r in data["leave"]:
            values = [r["name"], *[r["taken"][str(t["id"])] for t in types], r["total"], r["waiting"]]
            cells = [values[0], *[f"{v:g}" for v in values[1:]]]
            assert " ".join(reversed(cells) if lang == "ur" else cells) in lines
            assert tuple(got[r["name"]][: len(values)]) == tuple(values)
    colleague = next(r for r in data["leave"] if r["name"] == "Sara Colleague")
    assert colleague["total"] == 1.5
