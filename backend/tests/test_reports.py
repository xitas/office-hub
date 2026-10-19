"""Reports, slice 1: task completion and overdue tasks. The numbers (on time by the office day, a task with two
assignees counted once per person and once in the totals), who sees what (managers their department, admins all,
staff nothing; only visible tasks), periods, PDF and Excel exports (Urdu right to left with the Nastaliq font,
formula-safe cells, the export permission and audit log) and query counts that don't grow with the data."""
import io
from datetime import datetime, time, timedelta

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import AuditLog
from apps.core.office_time import office_today, office_tz
from apps.tasks.models import Task

pytestmark = pytest.mark.django_db

COMPLETION = "/api/v1/reports/task-completion/"
OVERDUE = "/api/v1/reports/overdue-tasks/"


@pytest.fixture
def people(make_user, admin, manager, staff):
    return {
        "admin": admin, "manager": manager, "staff": staff,
        "colleague": make_user("colleague@test.com", full_name="Sara Colleague"),
        "ops": make_user("ops.person@test.com", department="ops", full_name="Omar Ops"),
        "ops_manager": make_user("ops.manager@test.com", role="manager", department="ops", full_name="Olga Manager"),
    }


def at(day, hh, mm=0):
    return datetime.combine(day, time(hh, mm), tzinfo=office_tz())


def task(title, who, due, *, done_at=None, creator=None, priority="medium"):
    t = Task.objects.create(title=title, due_date=due, status="done" if done_at else "todo", completed_at=done_at, priority=priority,
                            created_by=creator or who[0])
    t.assignees.set(who)
    return t


@pytest.fixture
def world(people):
    today = office_today()
    d = today - timedelta(days=10)  # due in the period
    s, c, o = people["staff"], people["colleague"], people["ops"]
    task("On time, last minute", [s], d, done_at=at(d, 23, 30))  # 18:30 UTC: same office day
    task("Late, just after midnight", [s], d, done_at=at(d + timedelta(days=1), 0, 30))  # still the due day in UTC
    task("Overdue", [s, c], d)  # two assignees: once each, once in totals
    task("Still open", [s], today + timedelta(days=1))
    task("Done early, due later", [c], today + timedelta(days=3), done_at=at(today - timedelta(days=2), 10))
    task("Before the period", [s], today - timedelta(days=60), done_at=at(today - timedelta(days=5), 9))  # completed in it
    task("Other department", [o], d)
    task("Someone else's, not visible to staff", [c], d, creator=c)
    return {"first": (today - timedelta(days=20)).isoformat(), "last": (today + timedelta(days=5)).isoformat()}


def test_completion_numbers_by_person_department_and_overall(client_for, people, world):
    res = client_for(people["manager"]).get(COMPLETION, {"from": world["first"], "to": world["last"]})
    assert res.status_code == 200
    by = {p["name"]: p for p in res.data["people"]}
    assert set(by) == {"Manager", "Staff", "Sara Colleague"}  # the manager's department only
    s = by["Staff"]
    assert (s["due"], s["done"], s["on_time"], s["late"], s["overdue"], s["open"], s["completed"]) == (4, 2, 1, 1, 1, 1, 3)
    assert s["rate"] == 0.5 and s["on_time_rate"] == 0.25
    sara = by["Sara Colleague"]
    assert (sara["due"], sara["done"], sara["overdue"]) == (3, 1, 2)
    assert by["Manager"]["due"] == 0 and by["Manager"]["rate"] is None
    totals = res.data["totals"]
    assert (totals["due"], totals["done"], totals["overdue"]) == (6, 3, 2)  # "Overdue" counted once
    [sales] = res.data["departments"]
    assert sales["due"] == 6
    # Admins: everyone, or one department.
    admin = client_for(people["admin"])
    assert res.data["totals"]["due"] < admin.get(COMPLETION, {"from": world["first"], "to": world["last"]}).data["totals"]["due"]
    ops = admin.get(COMPLETION, {"from": world["first"], "to": world["last"], "department": people["ops"].department_id}).data
    assert [p["name"] for p in ops["people"] if p["due"]] == ["Omar Ops"]


def test_scope_permissions_and_periods(client_for, people, world):
    assert client_for(people["staff"]).get(COMPLETION).status_code == 403
    assert client_for(people["staff"]).get(OVERDUE).status_code == 403
    manager = client_for(people["manager"])
    # A manager can't widen the scope: another department is ignored, someone outside is refused.
    res = manager.get(COMPLETION, {"department": people["ops"].department_id})
    assert {p["name"] for p in res.data["people"]} == {"Manager", "Staff", "Sara Colleague"}
    assert manager.get(COMPLETION, {"person": people["ops"].pk}).status_code == 400
    one = manager.get(COMPLETION, {"person": people["staff"].pk, "from": world["first"], "to": world["last"]}).data
    assert [p["name"] for p in one["people"]] == ["Staff"] and one["totals"]["due"] == 4
    today = office_today()
    assert manager.get(COMPLETION, {"from": today.isoformat(), "to": (today - timedelta(days=1)).isoformat()}).status_code == 400
    assert manager.get(COMPLETION, {"from": (today - timedelta(days=400)).isoformat(), "to": today.isoformat()}).status_code == 400
    default = manager.get(COMPLETION).data
    assert default["from"] == today.replace(day=1) and default["to"] == today


def test_overdue_by_person_with_buckets(client_for, people):
    today = office_today()
    s = people["staff"]
    task("3 days late", [s], today - timedelta(days=3), priority="urgent")
    task("20 days late", [s], today - timedelta(days=20))
    task("45 days late", [s, people["colleague"]], today - timedelta(days=45))
    task("Due today", [s], today)
    task("Done late", [s], today - timedelta(days=9), done_at=timezone.now())
    res = client_for(people["manager"]).get(OVERDUE)
    assert res.status_code == 200
    staff = next(p for p in res.data["people"] if p["name"] == "Staff")
    assert (staff["overdue"], staff["days_1_7"], staff["days_8_30"], staff["days_31_plus"], staff["urgent"]) == (3, 1, 1, 1, 1)
    assert staff["oldest"] == today - timedelta(days=45)
    assert res.data["total"] == 3 and [t["title"] for t in res.data["tasks"]] == ["45 days late", "20 days late", "3 days late"]
    assert res.data["tasks"][0]["assignees"] == ["Sara Colleague", "Staff"] and res.data["tasks"][0]["days_late"] == 45


def test_exports_pdf_and_excel_in_both_languages(client_for, people, world):
    task("=HYPERLINK(\"http://x\")", [people["staff"]], office_today() - timedelta(days=2))  # starts like a formula
    manager = client_for(people["manager"])
    params = {"from": world["first"], "to": world["last"]}
    assert manager.get(f"{COMPLETION}export/", {**params, "file": "doc"}).status_code == 400
    pdf = manager.get(f"{COMPLETION}export/", {**params, "file": "pdf"})
    assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert b"NotoSans" in pdf.content
    xlsx = manager.get(f"{OVERDUE}export/", {"file": "xlsx"})
    assert xlsx.status_code == 200 and "attachment" in xlsx["Content-Disposition"]
    wb = load_workbook(io.BytesIO(xlsx.content))
    tasks = wb[wb.sheetnames[1]]
    cells = [c for row in tasks.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
    assert cells and all(c.data_type == "s" and c.quotePrefix for c in cells)  # text, not a formula
    assert not wb[wb.sheetnames[0]].sheet_view.rightToLeft
    # Urdu: the Nastaliq font in the PDF, right-to-left sheets.
    people["manager"].language = "ur"
    people["manager"].save(update_fields=["language"])
    ur = client_for(people["manager"])
    pdf_ur = ur.get(f"{COMPLETION}export/", {**params, "file": "pdf"}, HTTP_ACCEPT_LANGUAGE="ur")
    assert pdf_ur.status_code == 200 and b"NotoNastaliqUrdu" in pdf_ur.content
    wb_ur = load_workbook(io.BytesIO(ur.get(f"{COMPLETION}export/", {**params, "file": "xlsx"}, HTTP_ACCEPT_LANGUAGE="ur").content))
    assert wb_ur[wb_ur.sheetnames[0]].sheet_view.rightToLeft
    # Exports need the export permission and are audited; staff can't.
    assert client_for(people["staff"]).get(f"{OVERDUE}export/", {"file": "pdf"}).status_code == 403
    logged = AuditLog.objects.filter(action="export")
    assert logged.count() == 4 and set(logged.values_list("changes__records", flat=True)) == {"task completion report", "overdue tasks report"}


def test_report_queries_do_not_grow(client_for, people, make_user):
    admin = client_for(people["admin"])
    today = office_today()

    def count(url, params):
        admin.get(url, params)  # warm the settings cache
        with CaptureQueriesContext(connection) as ctx:
            res = admin.get(url, params)
        assert res.status_code == 200
        return len(ctx)

    def add(i):
        person = make_user(f"p{i}@test.com", department="ops" if i % 2 else "sales")
        for k in range(6):
            due = today - timedelta(days=k * 5)
            task(f"T{i}-{k}", [person, people["staff"]], due, done_at=at(due, 12) if k % 2 else None, priority="high" if k % 3 else "low")

    urls = {
        "completion": (COMPLETION, {"from": (today - timedelta(days=60)).isoformat(), "to": today.isoformat()}),
        "overdue": (OVERDUE, {}),
        "completion pdf": (f"{COMPLETION}export/", {"file": "pdf"}),
        "overdue xlsx": (f"{OVERDUE}export/", {"file": "xlsx"}),
        "people": ("/api/v1/reports/people/", {}),
    }
    add(0)
    small = {k: count(*v) for k, v in urls.items()}
    for i in range(1, 8):
        add(i)
    grown = {k: count(*v) for k, v in urls.items()}
    assert grown == small


def test_export_percentages_round_like_the_screen():
    from apps.reports.render import _text

    assert [_text(v, "pct") for v in (0.125, 0.135, 0.004, 1)] == ["13%", "14%", "0%", "100%"]  # half up, not to even
    assert _text(None, "pct") == "—"
