"""Reports, slice 4: attendance and leave summary. Days in the office, on field visits and remote, late, on leave and
absent over working days (holidays and days off left out, from the day each person joined, today not absent yet), the
same as the attendance page's day statuses; leave taken in the period (clipped, half days 0.5, waiting requests) and the
year's balances; by person and department. Scope per role, exports in both languages with the audit log, and query
counts that don't grow with the data."""
import io
from datetime import datetime, time, timedelta
from decimal import Decimal

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.attendance import status as att
from apps.attendance.models import Attendance
from apps.core.models import AuditLog
from apps.core.office_time import office_today, office_tz
from apps.events.models import Event
from apps.leave.models import LeaveRequest, LeaveType
from apps.reports.scope import day_bounds, scope_for

pytestmark = pytest.mark.django_db

URL = "/api/v1/reports/attendance/"


@pytest.fixture
def people(make_user, admin, manager, staff):
    out = {
        "admin": admin, "manager": manager, "staff": staff,
        "colleague": make_user("colleague@test.com", full_name="Sara Colleague"),
        "ops": make_user("ops.person@test.com", department="ops", full_name="Omar Ops"),
    }
    User.objects.update(date_joined=datetime(2020, 1, 1, tzinfo=office_tz()))
    for u in out.values():
        u.refresh_from_db()
    return out


@pytest.fixture
def types():
    return {
        "annual": LeaveType.objects.create(name="Annual", name_ur="سالانہ", days_per_year=Decimal("20"), position=1),
        "casual": LeaveType.objects.create(name="Casual", days_per_year=Decimal("8"), position=2),
        "unpaid": LeaveType.objects.create(name="Unpaid", days_per_year=None, paid=False, position=3),
    }


def past_week():
    """A Monday at least two weeks back whose week stays in one year."""
    monday = office_today() - timedelta(days=office_today().weekday() + 14)
    while (monday + timedelta(days=6)).year != monday.year or monday - timedelta(days=4) < monday.replace(month=1, day=1):
        monday -= timedelta(days=7)
    return monday


def check_in(user, day, hour, minute=0, mode="office"):
    Attendance.objects.create(user=user, date=day, mode=mode, check_in=datetime.combine(day, time(hour, minute), tzinfo=office_tz()))


def leave(user, kind, first, last, days, status="approved", half_day=""):
    LeaveRequest.objects.create(user=user, leave_type=kind, start_date=first, end_date=last, days=days, status=status, half_day=half_day)


def week(monday):
    return {"from": monday.isoformat(), "to": (monday + timedelta(days=6)).isoformat()}


def test_attendance_and_leave_summary(client_for, people, types, make_user):
    mon = past_week()
    day = lambda n: mon + timedelta(days=n)  # noqa: E731
    s, c = people["staff"], people["colleague"]
    start, end = day_bounds(day(2), day(2))
    Event.objects.create(kind="holiday", title="Holiday", all_day=True, start=start, end=end, organizer=people["admin"])
    check_in(s, day(0), 9)  # on time
    check_in(s, day(1), 9, 30)  # late (after 9:15)
    check_in(s, day(3), 10, mode="remote")  # late
    check_in(s, day(4), 11, mode="field")  # field visits are never late
    check_in(s, day(2), 10)  # the holiday: in the office, never late, not a working day
    check_in(s, day(5), 10)  # a Saturday: the same
    leave(c, types["annual"], day(-4), day(1), 4)  # Thursday before to Tuesday: 2 days in the period
    leave(c, types["casual"], day(3), day(3), Decimal("0.5"), half_day="morning")
    check_in(c, day(3), 13)  # half-day leave and checked in: present (and late)
    leave(c, types["unpaid"], day(4), day(4), 1, status="pending")  # waiting: Friday still absent
    newbie = make_user("new@test.com", full_name="New Person")
    User.objects.filter(pk=newbie.pk).update(date_joined=datetime.combine(day(3), time(8), tzinfo=office_tz()))
    check_in(people["ops"], day(0), 9)

    res = client_for(people["manager"]).get(URL, week(mon))
    assert res.status_code == 200
    rows = {p["name"]: p for p in res.data["people"]}
    assert set(rows) == {"Manager", "Staff", "Sara Colleague", "New Person"}
    pick = lambda r: tuple(r[k] for k in ("working_days", "office", "field", "remote", "late", "leave", "absent", "attended"))  # noqa: E731
    assert pick(rows["Staff"]) == (4, 4, 1, 1, 2, 0, 0, 4) and rows["Staff"]["rate"] == 1
    assert pick(rows["Sara Colleague"]) == (4, 1, 0, 0, 1, 2, 1, 1) and rows["Sara Colleague"]["rate"] == 0.5
    assert pick(rows["New Person"])[0::6] == (2, 2) and rows["New Person"]["rate"] == 0
    assert res.data["period_working_days"] == 4 and res.data["holidays"] == 1
    assert res.data["totals"]["absent"] == sum(r["absent"] for r in rows.values())
    assert [d["name"] for d in res.data["departments"]] == ["Sales"] and res.data["departments"][0]["people"] == 4

    leave_rows = {r["name"]: r for r in res.data["leave"]}
    t = {k: v.pk for k, v in types.items()}
    assert leave_rows["Sara Colleague"]["taken"] == {t["annual"]: 2, t["casual"]: 0.5, t["unpaid"]: 0}
    assert leave_rows["Sara Colleague"]["total"] == 2.5 and leave_rows["Sara Colleague"]["waiting"] == 1
    assert leave_rows["Sara Colleague"]["remaining"] == {t["annual"]: 16, t["casual"]: 7.5, t["unpaid"]: None}
    by_type = {x["name"]: x for x in res.data["leave_types"]}
    assert by_type["Annual"]["taken"] == 2 and by_type["Annual"]["people"] == 1 and by_type["Annual"]["allowance"] == 80
    assert by_type["Unpaid"]["allowance"] is None

    # The same days as the attendance page, person by person.
    for p in (s, c):
        infos = att.compute([p.pk], mon, day(6))
        statuses = [infos[(p.pk, d)]["status"] for d in (day(n) for n in range(7)) if d.weekday() < 5 and d != day(2)]
        assert statuses.count("absent") == rows[p.full_name]["absent"] and statuses.count("leave") == rows[p.full_name]["leave"]
        assert sum(i["late"] for i in infos.values()) == rows[p.full_name]["late"]

    # Admins: everyone, by department; staff can't see the report.
    admin = client_for(people["admin"]).get(URL, week(mon)).data
    assert [d["name"] for d in admin["departments"]] == ["Operations", "Sales", None]  # the admin has no department
    assert client_for(s).get(URL).status_code == 403


def test_today_is_not_absent_yet(client_for, people):
    today = office_today().isoformat()
    row = next(p for p in client_for(people["manager"]).get(URL, {"from": today, "to": today}).data["people"] if p["name"] == "Staff")
    assert row["absent"] == 0 and row["working_days"] in (0, 1)
    check_in(people["staff"], office_today(), 9)
    row = next(p for p in client_for(people["manager"]).get(URL, {"from": today, "to": today}).data["people"] if p["name"] == "Staff")
    assert row["office"] == 1 and row["absent"] == 0


def test_report_people_are_the_attendance_team(people, make_user):
    """The report's people are those whose attendance the viewer sees on the attendance page (its team views)."""
    from apps.attendance.views import team_users

    viewers = {**people, "solo": make_user("solo.manager@test.com", role="manager", department=None)}
    for who in ("manager", "admin"):
        assert {p.pk for p in scope_for(viewers[who], {}).people} == set(team_users(viewers[who]).values_list("pk", flat=True))
    assert [p.pk for p in scope_for(viewers["solo"], {}).people] == [viewers["solo"].pk]


def test_attendance_report_exports(client_for, people, types):
    mon = past_week()
    check_in(people["staff"], mon, 9, 40)
    leave(people["colleague"], types["annual"], mon, mon + timedelta(days=1), 2)
    leave(people["colleague"], types["casual"], mon + timedelta(days=3), mon + timedelta(days=3), Decimal("0.5"), half_day="morning")
    manager = client_for(people["manager"])
    for lang in ("en", "ur"):
        pdf = manager.get(f"{URL}export/", {**week(mon), "file": "pdf"}, HTTP_ACCEPT_LANGUAGE=lang)
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and (lang == "en" or b"NotoNastaliqUrdu" in pdf.content)
        wb = load_workbook(io.BytesIO(manager.get(f"{URL}export/", {**week(mon), "file": "xlsx"}, HTTP_ACCEPT_LANGUAGE=lang).content))
        assert all(ws.sheet_view.rightToLeft == (lang == "ur") for ws in wb.worksheets)
        taken = wb.worksheets[3]
        header = [c.value for c in taken[7]]
        assert header[1] == ("سالانہ" if lang == "ur" else "Annual")
        colleague = next(r for r in taken.iter_rows(min_row=8, values_only=True) if r[0] == "Sara Colleague")
        assert colleague[1:5] == (2, 0.5, 0, 2.5)
    assert client_for(people["staff"]).get(f"{URL}export/", {"file": "xlsx"}).status_code == 403
    assert set(AuditLog.objects.filter(action="export").values_list("changes__records", flat=True)) == {"attendance and leave report"}


def test_attendance_report_queries_do_not_grow(client_for, people, types, make_user):
    admin = client_for(people["admin"])
    mon = past_week()
    params = {"from": (mon - timedelta(days=60)).isoformat(), "to": (mon + timedelta(days=6)).isoformat()}

    def count(url, extra=None):
        admin.get(url, {**params, **(extra or {})})
        with CaptureQueriesContext(connection) as ctx:
            res = admin.get(url, {**params, **(extra or {})})
        assert res.status_code == 200
        return len(ctx)

    def add(i):
        person = make_user(f"p{i}@test.com", department="ops" if i % 2 else "sales")
        for k in range(5):
            check_in(person, mon - timedelta(days=7 * k), 9 + k % 2, mode=["office", "remote", "field"][k % 3])
        leave(person, types["annual" if i % 2 else "casual"], mon + timedelta(days=1), mon + timedelta(days=2), 2)
        leave(person, types["unpaid"], mon + timedelta(days=3), mon + timedelta(days=3), 1, status="pending")

    urls = {"report": (URL,), "pdf": (f"{URL}export/", {"file": "pdf"}), "xlsx": (f"{URL}export/", {"file": "xlsx"})}
    add(0)
    small = {k: count(*v) for k, v in urls.items()}
    for i in range(1, 7):
        add(i)
    assert {k: count(*v) for k, v in urls.items()} == small
